"""Importação: formatos, arquivos vazios/inválidos, PDFs de várias páginas e OCR."""

import io

import docx
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from formato_facil import extracao, fontes, imagens, importacao, ocr
from formato_facil.erros import ErroArquivo
from formato_facil.formatacao import ConfigTexto, calcular_layout
from formato_facil.geometria import pt_para_mm
from formato_facil.pdf_gerador import gerar_pdf_imagens, gerar_pdf_texto
from formato_facil.validacao import verificar_pdf

from .conftest import TEXTO_ACENTOS, imagem_bytes


def _pdf_com_paginas(textos):
    saida = io.BytesIO()
    c = canvas.Canvas(saida, pagesize=A4)
    interno = fontes.registrar_fonte(fontes.fonte_padrao())
    for texto in textos:
        if texto:
            c.setFont(interno, 14)
            c.drawString(72, 700, texto)
        else:  # página "digitalizada": só uma imagem, sem texto
            imagem = Image.new("RGB", (400, 200), "white")
            ImageDraw.Draw(imagem).rectangle([20, 20, 380, 180], outline="black", width=4)
            from reportlab.lib.utils import ImageReader
            c.drawImage(ImageReader(imagem), 72, 500, 300, 150)
        c.showPage()
    c.save()
    return saida.getvalue()


def _docx(paragrafos, tabela=None):
    documento = docx.Document()
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    if tabela:
        t = documento.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, valor in enumerate(linha):
                t.cell(i, j).text = valor
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


# --- TXT ---------------------------------------------------------------------

@pytest.mark.parametrize("codificacao", ["utf-8", "utf-8-sig", "cp1252", "utf-16"])
def test_txt_em_varias_codificacoes(codificacao):
    texto = "Ação e coração — “citação”."
    if codificacao == "cp1252":
        texto = "Ação e coração — “citação”."  # todos existem no Windows-1252
    documento = importacao.importar("nota.txt", texto.encode(codificacao))
    assert documento.texto == texto
    assert documento.tipo == "txt"


def test_txt_crlf_vira_lf():
    documento = importacao.importar("a.txt", "linha 1\r\nlinha 2\r\n\r\nlinha 3".encode())
    assert documento.texto == "linha 1\nlinha 2\n\nlinha 3"


# 6. Arquivos de texto vazios
def test_arquivo_vazio_gera_erro():
    with pytest.raises(ErroArquivo) as erro:
        importacao.importar("vazio.txt", b"")
    assert "vazio" in erro.value.mensagem


def test_arquivo_so_com_espacos_gera_aviso():
    documento = importacao.importar("branco.txt", b"   \r\n\t\n  ")
    assert documento.texto == ""
    assert any("não contém texto" in aviso for aviso in documento.avisos)


def test_docx_vazio_gera_aviso():
    documento = importacao.importar("vazio.docx", _docx([]))
    assert documento.texto == ""
    assert documento.avisos


# --- DOCX --------------------------------------------------------------------

def test_docx_paragrafos_e_tabela():
    dados = _docx(["Título do relatório", "", "Conteúdo com acentuação: ação."],
                  tabela=[["Nome", "Preço"], ["Pão", "R$ 1,50"]])
    documento = importacao.importar("rel.docx", dados)
    assert documento.texto == (
        "Título do relatório\n\nConteúdo com acentuação: ação.\n\nNome | Preço\n\nPão | R$ 1,50"
    )
    assert any("tabelas" in aviso for aviso in documento.avisos)


# --- PDF ---------------------------------------------------------------------

def test_pdf_com_texto_selecionavel():
    documento = importacao.importar("a.pdf", _pdf_com_paginas(["Olá, ação e coração."]))
    assert documento.texto == "Olá, ação e coração."
    assert not documento.precisa_ocr


# 10. PDFs de várias páginas são tratados de forma previsível
def test_pdf_varias_paginas_em_ordem():
    dados = _pdf_com_paginas(["Página um", "Página dois", "Página três"])
    documento = importacao.importar("tres.pdf", dados)
    assert documento.numero_paginas == 3
    assert documento.texto == "Página um\n\nPágina dois\n\nPágina três"
    assert any("3 páginas" in aviso for aviso in documento.avisos)


def test_pdf_paginas_sem_texto_sao_identificadas():
    dados = _pdf_com_paginas(["Página com texto", None, "Outra com texto"])
    documento = importacao.importar("misto.pdf", dados)
    assert documento.info_pdf.paginas_sem_texto == [1]
    assert documento.precisa_ocr
    assert any("2" in aviso and "OCR" in aviso for aviso in documento.avisos)


def test_pdf_varias_paginas_no_modo_aparencia(area_padrao):
    dados = _pdf_com_paginas(["Um", "Dois", "Três", "Quatro"])
    paginas = extracao.paginas_como_imagens(dados, area_padrao.largura_mm, area_padrao.altura_mm)
    pdf, _ = gerar_pdf_imagens(paginas, area_padrao, "ajustar")
    relatorio = verificar_pdf(pdf, area_padrao)
    assert len(relatorio.paginas) == 4
    assert relatorio.aprovado, relatorio.problemas
    caixas = {(round(p.caixa_imagens.x_mm, 3), round(p.caixa_imagens.y_mm, 3),
               round(p.caixa_imagens.largura_mm, 3)) for p in relatorio.paginas}
    assert len(caixas) == 1  # mesma posição em todas as folhas


def test_pdf_protegido_por_senha():
    documento = pymupdf.open()
    documento.new_page()
    dados = documento.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="dono", user_pw="senha")
    with pytest.raises(ErroArquivo) as erro:
        importacao.importar("secreto.pdf", dados)
    assert "senha" in erro.value.mensagem


# 9. Arquivos inválidos produzem mensagens de erro
@pytest.mark.parametrize("nome,dados,trecho", [
    ("programa.exe", b"MZ\x90\x00", "não é suportado"),
    ("sem_extensao", b"abc", "não é suportado"),
    ("falso.pdf", b"isto nao e um pdf", "não é um PDF válido"),
    ("quebrado.pdf", b"%PDF-1.4\n%%lixo", "PDF"),
    ("falso.docx", b"texto qualquer", "Word"),
    ("falso.png", b"isto nao e uma imagem", "imagem"),
    ("falso.jpg", b"\xff\xd8\xff\xe0quebrado", "imagem"),
    ("binario.txt", b"\x00\x01\x02\x00\x00", "não parece conter texto"),
])
def test_arquivos_invalidos(nome, dados, trecho):
    with pytest.raises(ErroArquivo) as erro:
        importacao.importar(nome, dados)
    assert trecho in erro.value.mensagem


def test_extensao_em_maiusculas():
    documento = importacao.importar("FOTO.JPG", imagem_bytes(100, 50, "JPEG"))
    assert documento.tipo == "imagem"


# --- Imagens -----------------------------------------------------------------

@pytest.mark.parametrize("formato,extensao", [("PNG", "png"), ("JPEG", "jpg"), ("WEBP", "webp")])
def test_formatos_de_imagem(formato, extensao):
    documento = importacao.importar(f"img.{extensao}", imagem_bytes(320, 200, formato))
    assert documento.imagem.size == (320, 200)
    assert documento.precisa_ocr and documento.permite_preservar_aparencia


def test_png_transparente_vira_fundo_branco():
    imagem = Image.new("RGBA", (10, 10), (0, 0, 0, 0))
    saida = io.BytesIO()
    imagem.save(saida, "PNG")
    carregada, _ = imagens.carregar_imagem(saida.getvalue())
    assert carregada.mode == "RGB" and carregada.getpixel((5, 5)) == (255, 255, 255)


# 8. Imagens com diferentes proporções são tratadas sem distorção
@pytest.mark.parametrize("largura,altura", [(1600, 400), (300, 2000), (1000, 1000), (750, 1400), (4000, 3000)])
def test_imagem_ajustada_sem_distorcao(area_padrao, largura, altura):
    imagem, _ = imagens.carregar_imagem(imagem_bytes(largura, altura))
    pdf, _ = gerar_pdf_imagens([imagem], area_padrao, "ajustar")
    with pymupdf.open(stream=pdf, filetype="pdf") as documento:
        caixa = pymupdf.Rect(documento[0].get_image_info()[0]["bbox"])
    proporcao_original = largura / altura
    assert caixa.width / caixa.height == pytest.approx(proporcao_original, rel=0.002)
    w, h = pt_para_mm(caixa.width), pt_para_mm(caixa.height)
    # Encosta em pelo menos um dos lados da área e não ultrapassa nenhum
    assert w <= 75.0 + 0.01 and h <= 140.0 + 0.01
    assert w == pytest.approx(75.0, abs=0.01) or h == pytest.approx(140.0, abs=0.01)
    assert verificar_pdf(pdf, area_padrao, tolerancia_mm=0.01).aprovado


@pytest.mark.parametrize("largura,altura", [(1600, 400), (300, 2000), (1000, 1000)])
def test_imagem_preenchida_corta_sem_distorcer(area_padrao, largura, altura):
    encaixe = imagens.calcular_encaixe(largura, altura, area_padrao, "preencher")
    esq, topo, dir_, base = encaixe.recorte_px
    assert (dir_ - esq) / (base - topo) == pytest.approx(75 / 140, rel=0.01)
    assert encaixe.fracao_imagem_cortada > 0
    assert imagens.explicar_proporcao(encaixe) is not None
    imagem, _ = imagens.carregar_imagem(imagem_bytes(largura, altura))
    pdf, _ = gerar_pdf_imagens([imagem], area_padrao, "preencher")
    with pymupdf.open(stream=pdf, filetype="pdf") as documento:
        info = documento[0].get_image_info()[0]
        caixa = pymupdf.Rect(info["bbox"])
        # A imagem embutida tem a mesma proporção da caixa onde foi desenhada
        assert info["width"] / info["height"] == pytest.approx(caixa.width / caixa.height, rel=0.01)
    assert pt_para_mm(caixa.width) == pytest.approx(75, abs=0.01)
    assert pt_para_mm(caixa.height) == pytest.approx(140, abs=0.01)


def test_explicacao_de_proporcao_no_modo_ajustar(area_padrao):
    encaixe = imagens.calcular_encaixe(1600, 900, area_padrao, "ajustar")
    mensagem = imagens.explicar_proporcao(encaixe)
    assert "espaço em branco" in mensagem and "acima e abaixo" in mensagem
    igual = imagens.calcular_encaixe(750, 1400, area_padrao, "ajustar")
    assert imagens.explicar_proporcao(igual) is None


def test_imagem_de_baixa_resolucao_gera_aviso(area_padrao):
    imagem, _ = imagens.carregar_imagem(imagem_bytes(100, 180))
    _, avisos = gerar_pdf_imagens([imagem], area_padrao, "ajustar")
    assert any("resolução baixa" in aviso for aviso in avisos)


# --- OCR ---------------------------------------------------------------------

def _imagem_com_texto(linhas):
    caminho = fontes.fontes_disponiveis()[fontes.fonte_padrao()]
    fonte = ImageFont.truetype(str(caminho), 42)
    imagem = Image.new("RGB", (1400, 120 + 70 * len(linhas)), "white")
    desenho = ImageDraw.Draw(imagem)
    for i, linha in enumerate(linhas):
        desenho.text((60, 60 + 70 * i), linha, font=fonte, fill="black")
    return imagem


requer_ocr = pytest.mark.skipif(not ocr.verificar_ocr().disponivel, reason="Tesseract não instalado")


@requer_ocr
def test_ocr_em_portugues():
    status = ocr.verificar_ocr()
    if (ocr.PASTA_TESSDATA_PROJETO / "por.traineddata").is_file():
        # O idioma português do aplicativo precisa ser encontrado (mesmo com
        # acentos no caminho da pasta, como em "formatação").
        assert status.tem_portugues, status.mensagem
    if not status.tem_portugues:
        pytest.skip("Idioma português do OCR não instalado")
    linhas = ["A educação é a base da nação.", "Coração, ação e informação."]
    resultado = ocr.reconhecer_texto(_imagem_com_texto(linhas))
    texto = " ".join(resultado.texto.split())
    assert resultado.idioma == "por"
    assert "educação" in texto and "nação" in texto and "Coração" in texto
    assert resultado.confianca_media > 75


@requer_ocr
def test_ocr_imagem_em_branco_avisa():
    resultado = ocr.reconhecer_texto(Image.new("RGB", (800, 400), "white"))
    assert resultado.texto.strip() == ""
    assert resultado.avisos


@requer_ocr
def test_ocr_em_pdf_digitalizado():
    imagem = _imagem_com_texto(["Documento digitalizado de teste"])
    saida = io.BytesIO()
    imagem.save(saida, "PDF", resolution=150)
    documento = importacao.importar("scan.pdf", saida.getvalue())
    assert not documento.info_pdf.tem_texto and documento.precisa_ocr
    resultado = importacao.executar_ocr(documento)
    assert "digitalizado" in resultado.texto


@requer_ocr
def test_fluxo_completo_imagem_ocr_para_pdf(area_padrao):
    """Modo B, opção 2: OCR de um print reconstruído como texto na área."""
    dados = io.BytesIO()
    _imagem_com_texto(["Lembrete: reunião às 14h", "na sala de reuniões."]).save(dados, "PNG")
    documento = importacao.importar("print.png", dados.getvalue())
    resultado = importacao.executar_ocr(documento)
    layout = calcular_layout(resultado.texto, ConfigTexto(), area_padrao)
    pdf = gerar_pdf_texto(layout)
    relatorio = verificar_pdf(pdf, area_padrao)
    assert relatorio.aprovado
    if ocr.verificar_ocr().tem_portugues:
        assert "reunião" in relatorio.texto_extraido


def test_ocr_indisponivel_gera_erro_claro(monkeypatch):
    from formato_facil.erros import ErroOCR
    monkeypatch.setattr(ocr, "verificar_ocr", lambda: ocr.StatusOCR(False, mensagem="Tesseract ausente."))
    with pytest.raises(ErroOCR) as erro:
        ocr.reconhecer_texto(Image.new("RGB", (10, 10)))
    assert "Tesseract" in erro.value.mensagem


def test_acentos_no_texto_gerado_a_partir_de_txt(area_padrao):
    documento = importacao.importar("acentos.txt", TEXTO_ACENTOS.encode("utf-8"))
    layout = calcular_layout(documento.texto, ConfigTexto(), area_padrao)
    relatorio = verificar_pdf(gerar_pdf_texto(layout), area_padrao)
    assert " ".join(relatorio.texto_extraido.split()) == " ".join(TEXTO_ACENTOS.split())
